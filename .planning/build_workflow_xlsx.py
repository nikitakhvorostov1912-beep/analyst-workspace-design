# -*- coding: utf-8 -*-
"""Генератор операционного workflow Excel для проекта «1С Аналитик».

Источник истины: .planning/ROADMAP-2026-05-29.md + knowledge-layer phases/M-Kx/.
Пересборка:  <system python> .planning/build_workflow_xlsx.py
Выход:       .planning/Workflow_1C_Analyst_2026-05-29.xlsx
"""
from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ── Брендовая палитра (Stencil/Mono) ────────────────────────────────────────
INK = "15161A"        # тёмный фон шапки
SIGNAL = "FF6A3D"     # оранжевый акцент
TINT = "FFE3D5"       # светлый оранжевый (заголовки секций)
WHITE = "FFFFFF"
GREEN = "C9F2D6"      # готово
AMBER = "FCE9C2"      # в работе / pending
GRAY = "ECECEC"       # не начато
ROW_ALT = "F7F6F3"    # чередование (Sand-tint)
P0 = "F6C5BE"         # приоритет P0
P1 = "FBE2BD"
P2 = "E6E6E6"

THIN = Side(style="thin", color="D9D6CF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HEAD_FONT = Font(name="Calibri", bold=True, color=WHITE, size=11)
TITLE_FONT = Font(name="Calibri", bold=True, color=WHITE, size=14)
CELL_FONT = Font(name="Calibri", size=10)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def status_fill(v: str) -> str | None:
    if v.startswith("✅"):
        return GREEN
    if v.startswith("🟡"):
        return AMBER
    if v.startswith("⬜"):
        return GRAY
    return None


def prio_fill(v: str) -> str | None:
    return {"P0": P0, "P1": P1, "P2": P2}.get(v.strip())


def add_sheet(wb, name, title, columns, rows, status_col=None, prio_col=None):
    """columns: list of (header, width, wrap). rows: list of tuples."""
    ws = wb.create_sheet(name)
    ncol = len(columns)

    # Заголовок-плашка (брендовая)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    tc = ws.cell(row=1, column=1, value=title)
    tc.font = TITLE_FONT
    tc.fill = PatternFill("solid", fgColor=INK)
    tc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 26

    # Шапка таблицы
    hr = 2
    for c, (header, width, _wrap) in enumerate(columns, start=1):
        cell = ws.cell(row=hr, column=c, value=header)
        cell.font = HEAD_FONT
        cell.fill = PatternFill("solid", fgColor=SIGNAL)
        cell.alignment = CENTER
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[hr].height = 30

    # Данные
    for r, row in enumerate(rows, start=hr + 1):
        alt = (r - hr) % 2 == 0
        for c, (header, _w, wrap) in enumerate(columns, start=1):
            val = row[c - 1] if c - 1 < len(row) else ""
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = CELL_FONT
            cell.border = BORDER
            cell.alignment = WRAP_TOP if wrap else CENTER
            fill = None
            if status_col is not None and c == status_col:
                fill = status_fill(str(val))
            elif prio_col is not None and c == prio_col:
                fill = prio_fill(str(val))
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
            elif alt:
                cell.fill = PatternFill("solid", fgColor=ROW_ALT)

    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(ncol)}{hr + len(rows)}"
    ws.sheet_view.showGridLines = False
    return ws


# ════════════════════════════════════════════════════════════════════════════
# Лист 1 — ОБЗОР (милстоуны)
# ════════════════════════════════════════════════════════════════════════════
OVERVIEW_COLS = [
    ("Этап", 10, False), ("Название", 30, True), ("Цель (одной фразой)", 46, True),
    ("Уровни L", 12, True), ("Бизнес-триггер", 30, True),
    ("Усилие", 12, True), ("Зависит от", 12, True), ("Статус", 22, True),
]
OVERVIEW = [
    ("Оболочка", "Продукт M1–M11 + Hermes + Commerce", "Чат-UI, SSE, 6 карточек, сессии, память/скиллы, Electron-инсталлятор, security", "L0–L1", "Платформа доставки", "—", "—", "✅ Готово (код v1.4.8)"),
    ("M-K0", "Stabilization", "Закрыть 28 critical+high из аудита перед Knowledge Layer", "—", "Фундамент стабилен", "~3-4 нед*", "—", "✅ Готово (28/28)"),
    ("M-K1", "Foundation + Multi-MCP", "Архрешения (ADR), capability discovery, multi-MCP, паспорт объекта", "L1", "Quick wins", "~2-3 нед*", "M-K0", "✅ Готово (15/17)"),
    ("M-K2", "Knowledge Foundation + Triple RAG", "Indexer + vector store + RAG ИТС (317) + RAG БСП (~3000 методов)", "L3, L5", "«Мозг» знает стандарты", "~4-5 нед*", "M-K1", "✅ Готово (12/13)"),
    ("M-K2.5", "Перестройка карточек (NVIDIA NIM)", "63 203 эталонных карточки типовых через qwen-122b", "L3/L5 enrich", "Качество ответов по типовым", "фон", "M-K2", "🟡 ~6.5% (фон, не блокер)"),
    ("M-K3", "Relational + Behavioral + EPF/CFE", "L2-граф + diagnose + EPF/CFE доставка + top-5 UC", "L2, L4", "MVP «Объяснитель» — killer USP", "~48 дн", "M-K2", "🟡 ~5% (EPF-скелет)"),
    ("M-K4", "Visual + Activity Stream + Reasoning", "Граф вызовов + Activity Stream + Posting Trace + L4 reasoning + hybrid", "L4, L5", "INFOSTART A&PM EVENT окт-2026", "~42-55 дн", "M-K3", "🟡 часть авансом"),
    ("M-K5", "Predictive + 8.5-Ready + Installer v2", "Temporal/сравнение с типовой + 8.5-Ready Assessment (доход) + installer v2.0", "L6", "ПЕРВЫЙ ПРЯМОЙ ДОХОД", "~5-7 нед", "M-K3,K4", "⬜ Не начато"),
    ("M-K6", "Predictive++ + Enterprise", "NL2SQL + cross-config benchmark + auto code review + pricing/франшиза", "L6+", "Moat, MRR ≥500к₽", "~6-10 нед", "все", "⬜ Не начато"),
]
OVERVIEW_NOTE = "* Календарные оценки старого плана (на команду). Соло+AI сжал M-K0→M-K2 из ~11 недель в дни. Планировать ПОРЯДКОМ и УСИЛИЕМ, не календарём. Жёсткие — только внешние якоря (см. лист «Календарь и якоря»)."


# ════════════════════════════════════════════════════════════════════════════
# Лист 2 — WORKFLOW (детально: что и КАК делается)
# ════════════════════════════════════════════════════════════════════════════
WF_COLS = [
    ("ID", 9, False), ("Этап", 9, False), ("Фаза", 16, True), ("Задача", 30, True),
    ("Что делается", 40, True), ("Как (метод / инструмент)", 40, True),
    ("Артефакт", 26, True), ("DoD / критерий", 34, True),
    ("Зависит", 12, True), ("Дн", 6, False), ("Приор.", 8, False), ("Статус", 18, True),
]

# Каждая строка: ID, Этап, Фаза, Задача, Что, Как, Артефакт, DoD, Завис, Дн, Приор, Статус
WORKFLOW = [
    # ── Немедленные действия D0–D3 ──
    ("D0", "Гигиена", "Док-синхро", "Свести 3 STATE к одному факту",
     "Привести knowledge-layer/STATE.md, .planning/STATE.md, CLAUDE.md-snapshot к единому состоянию; старые роадмапы → ссылка на ROADMAP-2026-05-29",
     "Edit (без удаления — история сохраняется)", "STATE ×3 + пометки", "Один факт состояния во всех доках", "—", "0.5", "P0", "⬜ Не начато"),
    ("D1", "M-K4(ав.)", "Коммит аванса", "Закоммитить guardrail G2",
     "Атомарный коммит уже протестированной интеграции bsp_warning в loop",
     "git add events.py loop.py test_loop_bsp_warning.py; commit -F", "commit", "2/2 теста зелёные в истории", "A-4", "0.2", "P0", "🟡 Готов к коммиту"),
    ("D1b", "M-K4(ав.)", "Frontend", "Фронт-баннер bsp_warning",
     "Обработать SSE-событие bsp_warning → неблокирующий баннер «сигнатуры БСП не верифицированы»",
     "React, useChatStream hook, vitest", "компонент + тесты", "Баннер появляется на phantom-ответе", "D1", "1", "P1", "⬜ Не начато"),
    ("D3", "M-K2.5", "Решение по NIM", "Судьба rebuild-карточек",
     "Проверить жив ли NIM-процесс; зафиксировать достигнутое или снять с критического внимания (enrichment, не блокер)",
     "watchdog-скрипт, проверка app.db/checkpoint", "решение в STATE", "Статус трека однозначен", "—", "0.5", "P1", "⬜ Не начато"),

    # ── M-K3.0 Preparatory ──
    ("K3.0a", "M-K3", "3.0 Prep", "Декомпозиция loop.py (ARCH-1)",
     "Выделить helpers из оркестратора (сейчас ~1835 строк), ядро ≤400 строк — чистота под интеграцию Knowledge Layer",
     "Рефакторинг, pytest-регрессия после каждого шага", "loop.py + новые модули", "Ядро ≤400 строк, все тесты зелёные", "—", "5", "P0", "⬜ Не начато"),
    ("K3.0b", "M-K3", "3.0 Prep", "Собрать ≥20 diagnose-кейсов",
     "Реальные кейсы (RLS-блокировка / пустой отчёт / медленный запрос) из практики Транзит/УСО — разблокируют UC 17.2–17.6 (test-first)",
     "Ручной сбор → golden dataset (UC = запрос+ожидаемый ответ)", "golden_diagnose.jsonl", "≥20 кейсов с ожидаемым ответом", "—", "2", "P0", "⬜ Не начато"),

    # ── Phase 13a EPF ──
    ("13a.0", "M-K3", "13a EPF", "EPF skeleton + epf-init",
     "Каркас внешней обработки АналитикLite",
     "skill epf-*, EDT", "АналитикLite.xml / ObjectModule.bsl", "epf-init собирается", "M-K2", "0.5", "P0", "✅ Готово (0a3b702)"),
    ("13a.1", "M-K3", "13a EPF", "manifest + manifest.json",
     "Метаданные обработки + json-манифест",
     "EDT, UTF-8 BOM", "manifest.json", "Манифест валиден", "13a.0", "0.5", "P0", "✅ Готово (0a3b702)"),
    ("13a.2", "M-K3", "13a EPF", "Capability response (8 caps)",
     "ADR-004: ответ о возможностях канала (8 базовых capabilities)",
     "BSL, JSON-схема experimental.analyst-1c.features", "capability endpoint", "Возвращает 8 base caps", "13a.1", "1.5", "P0", "⬜ Не начато"),
    ("13a.3", "M-K3", "13a EPF", "Общие модули АП_*",
     "HMAC-подпись + Capability (префикс АП_ по решению Q1)",
     "BSL общие модули", "АП_HMAC, АП_Capability", "HMAC-подпись проходит проверку", "13a.2", "2", "P0", "⬜ Не начато"),
    ("13a.4", "M-K3", "13a EPF", "HTTP-сервис :6011 + initialize",
     "EPF поднимает локальный HTTP-сервис",
     "BSL HTTP-сервис 1С", "сервис на :6011", "initialize отвечает корректно", "13a.3", "2", "P0", "⬜ Не начато"),
    ("13a.5", "M-K3", "13a EPF", "Onboarding step 1.5 «Скачать EPF»",
     "Шаг онбординга: выбор и скачивание EPF",
     "Next.js, OnboardingDialog", "шаг онбординга", "Шаг live", "13a.4", "1", "P1", "⬜ Не начато"),
    ("13a.6", "M-K3", "13a EPF", "Smoke EPF на УТ 11.5",
     "Проверка EPF на типовой УТ 11.5",
     "стенд 1С УТ 11.5", "smoke-результат", "EPF работает на УТ", "13a.4", "1", "P0", "⬜ Не начато"),
    ("13a.7", "M-K3", "13a EPF", "Smoke EPF на ERP 2.5", "Проверка на ERP", "стенд ERP 2.5", "smoke", "Работает", "13a.6", "0.5", "P1", "⬜ Не начато"),
    ("13a.8", "M-K3", "13a EPF", "Smoke EPF на КА 2.5", "Проверка на КА", "стенд КА 2.5", "smoke", "Работает", "13a.6", "0.5", "P1", "⬜ Не начато"),
    ("13a.9", "M-K3", "13a EPF", "EPF installer (PowerShell)",
     "Скрипт установки EPF + ручная инструкция",
     "PowerShell", "Install-АналитикLite.ps1", "Ставится на чистой 1С", "13a.6", "1", "P1", "⬜ Не начато"),
    ("13a.10", "M-K3", "13a EPF", "EPF release notes", "Документация по EPF", "markdown", "release notes", "Написаны", "13a.9", "0.5", "P2", "⬜ Не начато"),

    # ── Phase 13b CFE ──
    ("13b.0", "M-K3", "13b CFE", "feature-detection per config",
     "Определение доступных фич на каждой конфигурации (R-04 mitigation для БГУ/ЗУП)",
     "BSL метаданные-проба", "feature-map", "Корректно для 5 типовых", "13a.*", "0.5", "P0", "⬜ Не начато"),
    ("13b.1.1", "M-K3", "13b CFE", "CFE skeleton + cfe-init", "Каркас расширения конфигурации", "skill cfe-*, EDT", "АналитикПлюс.cfe скелет", "cfe-init собирается", "13b.0", "0.5", "P0", "⬜ Не начато"),
    ("13b.1.2", "M-K3", "13b CFE", "Подсистема АналитикПлюс", "Подсистема + префикс АП_ (Q1)", "EDT", "подсистема", "Создана", "13b.1.1", "0.5", "P0", "⬜ Не начато"),
    ("13b.1.3", "M-K3", "13b CFE", "Заимствование 3 общих модулей", "Перенос HMAC/Capability из EPF в CFE", "EDT заимствование", "модули в CFE", "Модули работают в CFE", "13b.1.2", "1", "P0", "⬜ Не начато"),
    ("13b.2", "M-K3", "13b CFE", "Activity Stream подписки",
     "Подписки на события (ПриЗаписи/ПриУдалении) — данные для визуализации в M-K4",
     "BSL подписки на события", "обработчики событий", "События ловятся", "13b.1.3", "2", "P1", "⬜ Не начато"),
    ("13b.3", "M-K3", "13b CFE", "Posting Trace (ПередЗаписью)",
     "Трассировка проведения документов — данные для M-K4 «почему записался»",
     "BSL ПередЗаписью, method override", "trace-данные", "Trace пишется", "13b.1.3", "2", "P1", "⬜ Не начато"),
    ("13b.4", "M-K3", "13b CFE", "BSL LS diagnostics через CFE",
     "Прокидывание диагностик BSL LS через CFE (detector-режим, G5)",
     "BSL + subprocess BSL LS", "diagnostics-канал", "Диагностики доходят", "13b.1.3", "3", "P1", "⬜ Не начато"),
    ("13b.4.1", "M-K3", "13b CFE", "HMAC SSO между EPF и CFE",
     "Единый вход 1С-пользователя через HMAC",
     "BSL HMAC", "SSO", "Сессия общая EPF↔CFE", "13b.4", "1", "P1", "⬜ Не начато"),
    ("13b.5", "M-K3", "13b CFE", "Capability extended (23 caps)",
     "Полный список возможностей CFE (ADR-004 full)",
     "BSL, JSON-схема", "capability extended", "Возвращает 23 caps", "13b.4.1", "1", "P0", "⬜ Не начато"),
    ("13b.5.1", "M-K3", "13b CFE", "Migration EPF→CFE",
     "Миграция без потери истории сессий при переходе EPF→CFE",
     "backend migration + проверка", "migration", "История сессий сохранена", "13b.5", "1.5", "P0", "⬜ Не начато"),
    ("13b.5.2", "M-K3", "13b CFE", "Smoke на 5 типовых",
     "Проверка CFE на УТ/ERP/КА/БГУ/ЗУП (Q2)",
     "5 стендов 1С", "5 smoke-наборов", "Применяется на всех 5", "13b.5.1", "3", "P0", "⬜ Не начато"),
    ("13b.6", "M-K3", "13b CFE", "CFE installer (signed PS)",
     "PowerShell-инсталлятор CFE, подписанный (или self-signed README)",
     "PowerShell, EV/OV cert (DEVOPS-1)", "Install-АналитикПлюс.ps1", "Ставится; подпись/README", "13b.5.2", "1.5", "P1", "⬜ Не начато"),

    # ── Phase 15 BSL LS Detector ──
    ("15.1", "M-K3", "15 BSL LS", "BSL LS jar + JRE bundle",
     "Поставка BSL LS jar + JRE (R-05: jlink ~30-40МБ + proguard ~50МБ, иначе download-on-demand)",
     "jlink, proguard, bsl-language-server-0.29.0.jar", "bundled jar+jre", "Запускается без системной Java", "—", "1", "P1", "⬜ Не начато"),
    ("15.2", "M-K3", "15 BSL LS", "Backend subprocess wrapper",
     "Обёртка запуска BSL LS как subprocess из backend",
     "Python subprocess, async", "bsl_ls_runner.py", "Возвращает диагностики JSON", "15.1", "1", "P1", "⬜ Не начато"),
    ("15.3", "M-K3", "15 BSL LS", "Antipattern detector pipeline",
     "Пайплайн прогона кода через детектор",
     "Python pipeline", "detector pipeline", "Прогоняет файл → список нарушений", "15.2", "1", "P1", "⬜ Не начато"),
    ("15.4", "M-K3", "15 BSL LS", "A1–A11 rules adapter",
     "Адаптер под AI-антипаттерны A1–A11",
     "rules/1c/1c-anti-patterns.md → adapter", "rules adapter", "Ловит A1–A11", "15.3", "1", "P1", "⬜ Не начато"),
    ("15.5", "M-K3", "15 BSL LS", "BSLDiagnostics card",
     "Карточка inline-warnings в чате",
     "React, card-registry", "BSLDiagnostics card", "Рендерит warnings", "15.4", "1", "P1", "⬜ Не начато"),

    # ── Phase 17 Top-5 UC ──
    ("17.1", "M-K3", "17 UC", "Knowledge Graph (КРИТ. ПУТЬ)",
     "L2-граф: AST → узлы/рёбра. БЕЗ него L4-diagnose невозможен",
     "TreeSitter BSL + SQLite + recursive CTE (ADR-002)", "graph.sqlite + builder", "~5K nodes/~50K edges, traversal 5 ур. ≤300мс; парсер ≥95% БСП 3.1", "K3.0a", "3", "P0", "⬜ Не начато"),
    ("17.2", "M-K3", "17 UC", "UC: RLS-tracer",
     "«Почему Иванов не видит документ» → trace до RLS-условия",
     "graph + MCP get_access_rights + diagnose", "RLS-tracer UC", "Точность ≥85% на golden (20 кейсов)", "17.1,K3.0b", "2", "P0", "⬜ Не начато"),
    ("17.3", "M-K3", "17 UC", "UC: Report-tracer",
     "«Почему отчёт пустой» → анализ СКД/параметров/фильтров",
     "graph + СКД-анализ + diagnose", "Report-tracer UC", "Pinpoint антипаттерна", "17.1,K3.0b", "2", "P0", "⬜ Не начато"),
    ("17.4", "M-K3", "17 UC", "UC: Цепочка вызовов",
     "«Покажи цепочку проведения документа» → движения в регистры",
     "call graph traversal (CTE)", "Lineage UC + GraphCard", "Граф движений строится", "17.1", "2", "P0", "⬜ Не начато"),
    ("17.5", "M-K3", "17 UC", "UC: Impact analysis",
     "«Что сломается если переименую реквизит X» → ≥5 мест с file:line",
     "reverse call graph + find_references", "Impact UC", "≥5 мест с file:line", "17.1", "2", "P0", "⬜ Не начато"),
    ("17.6", "M-K3", "17 UC", "UC: Query optimizer",
     "«Почему запрос медленный» → исправленная версия + объяснение",
     "antipattern detector + graph + LLM CoT", "Query-optimizer UC", "Исправленная версия + объяснение", "17.1,15.4", "2", "P1", "⬜ Не начато"),
    ("17.7", "M-K3", "17 UC", "GraphCard + DiagnoseCard + registry",
     "Визуальные карточки + типизированный card-registry на 19 типов (G11)",
     "React Flow, frontend/lib/card-registry.ts", "GraphCard, DiagnoseCard, registry", "Карточки рендерятся, registry типизирован", "17.4", "2", "P1", "⬜ Не начато"),
    ("K3.99", "M-K3", "SUMMARY", "M-K3 SUMMARY + handoff",
     "Итоговый документ + pre-flight для M-K4",
     "markdown по образцу M-K0/SUMMARY.md", "M-K3/SUMMARY.md", "Все фазы закрыты с commit hashes", "17.7,13b.6", "1", "P0", "⬜ Не начато"),

    # ── M-K4 (часть авансом + основное) ──
    ("A-1", "M-K4", "Retrieval (аванс)", "Hybrid retrieval (FTS5+RRF)",
     "BM25 (FTS5) ⊕ vector → RRF-merge поверх its/bsp search; работает без ключа на готовом индексе",
     "SQLite FTS5, RRF k=60, sqlite-vec", "hybrid.py, its/bsp_search.py", "Precision выше vector-only", "M-K2", "2", "P1", "✅ Готово (581fec9)"),
    ("A-2", "M-K4", "Guardrails (аванс)", "Phantom-method detector",
     "Извлечь BSL-вызовы из ответа, сверить с bsp_chunks → выдуманные методы",
     "regex, aiosqlite", "phantom_check.py", "Модуль известен + метод отсутствует → phantom", "M-K2", "1", "P1", "✅ Готово (aff4fb2)"),
    ("A-3", "M-K4", "Eval (аванс)", "Golden dataset + phantom-eval",
     "20 кейсов на РЕАЛЬНЫХ сигнатурах БСП; precision/recall/f1",
     "golden.jsonl, eval CLI", "its_kb_eval.py, golden.jsonl", "Eval считает TP/FP/FN", "M-K2", "1", "P1", "✅ Готово (049680b)"),
    ("A-4", "M-K4", "Guardrails (аванс)", "G2/G3 + интеграция в loop",
     "bsp_warning SSE (неблокирующий), G3 unsupported claim",
     "Pydantic event, loop hook, pytest", "guardrails.py, events.py, loop.py", "G2 эмитит bsp_warning, не блокирует done (2/2)", "A-2", "1", "P0", "🟡 Готов к коммиту (D1)"),
    ("16.0", "M-K4", "16 MetaVision", "MetaVision Spike (HARD GATE)",
     "go/no-go за 3 дня: headless CLI без X11, JSON nodes/edges, ≤60с на УТ 11.5",
     "MetaVisionFor1C jar, Java subprocess", "go/no-go решение", "Решение принято в timebox", "M-K3", "3", "P1", "⬜ Не начато"),
    ("16.x", "M-K4", "16 Visual", "MetaVisionGraph ИЛИ наш GraphCard",
     "GO → MetaVisionGraph (D3.js); NO-GO → наш GraphCard primary (6 дн вместо 15)",
     "D3.js / sigma.js / React Flow", "Graph card", "Граф вызовов на 1+ типовой", "16.0", "6-15", "P1", "⬜ Не начато"),
    ("16.A", "M-K4", "16.A Activity", "Activity Stream sidebar",
     "Real-time поток событий 1С из CFE + агрегация + TimelineCard",
     "SSE endpoint, debounce/batching, React sidebar", "Activity sidebar, TimelineCard", "События из CFE видны live", "M-K3 13b.2", "10", "P1", "⬜ Не начато"),
    ("16.B", "M-K4", "16.B Posting", "Posting Trace card + L4-связь",
     "Карточка трассировки проведения + связь с L4 «почему документ записался»",
     "knowledge.db trace storage, React", "PostingTrace card", "Trace отображается + объяснение", "M-K3 13b.3", "8", "P1", "⬜ Не начато"),
    ("17.A", "M-K4", "17.A Reasoning", "L4 Behavioral Reasoning",
     "Diagnose Engine YAML-rulebook + hypothesis chain + deadlock-tracer (L4-4) + compliance (L5-4) + refactor planner (L5-5)",
     "YAML rulebook, LLM CoT, hypothesis framework", "diagnose engine + DiagnoseCard/ProcessCard", "Отвечает «почему» с пруфами; compliance PDF", "M-K3 17.*", "12", "P1", "⬜ Не начато"),
    ("X-5b", "M-K4", "Retrieval", "Reranker поверх RRF",
     "Углубление аванса: reranker для precision@5 ≥85%",
     "cross-encoder / local reranker", "reranker", "precision@5 ≥85%", "A-1", "2", "P2", "⬜ Не начато"),
    ("UX-7", "M-K4", "Citation", "Citation Layer (G1)",
     "Каждое утверждение LLM с источником + hover-preview (как Perplexity). RAGResponse citation-контракт — ИНВАЗИВНО, за feature-flag",
     "RAGResponse контракт, feature-flag, React hover", "citation layer", "Каждое утверждение с источником", "A-1", "3", "P2", "⬜ Не начато"),
    ("OPS-2", "M-K4", "Telemetry", "Retrieval source telemetry",
     "Метрики источника ответа + hit-rate — ДАННЫЕ для build-vs-buy по Напарнику (собрать до 01.10.2026)",
     "sqlite metrics, sparkline dashboard", "telemetry + Metrics card", "Дашборд по источникам", "M-K2", "3", "P1", "⬜ Не начато"),
    ("OPS-5", "M-K4", "MCP server", "Knowledge Layer как MCP :7070",
     "Экспонировать Knowledge Layer отдельным MCP-сервером — чистая граница + сторонние ИИ",
     "MCP Streamable HTTP server", "MCP :7070", "Подключается как MCP-tool", "17.A", "2", "P2", "⬜ Не начато"),
    ("15.A", "M-K4", "15.A (опц.)", "BSL LS streaming",
     "WebSocket live-highlighting (опционально, если цикл позволит)",
     "WebSocket, debounce 500мс", "live diagnostics", "Live-подсветка warnings", "M-K3 15.5", "5", "P2", "⬜ Не начато"),
    ("K4.99", "M-K4", "SUMMARY", "M-K4 SUMMARY + handoff", "Итог + pre-flight M-K5", "markdown", "M-K4/SUMMARY.md", "Фазы закрыты", "17.A", "1", "P1", "⬜ Не начато"),

    # ── M-K5 (крупно) ──
    ("L6-3", "M-K5", "Reference lib", "Reference Configurations Library",
     "Fingerprints типовых УТ 11.5/ERP 2.5/БП 3.0/УСО 2.5 (только метаданные+структура, НЕ код — юр.риск)",
     "fingerprint + структура в SQLite", "reference library", "4 типовые зафиксированы", "M-K3", "—", "P1", "⬜ Не начато"),
    ("L6-2", "M-K5", "Compare", "UC: Сравни с типовой + ComparisonCard",
     "Отчёт об отличиях клиентской базы от типовой",
     "diff fingerprint + graph", "Compare UC, ComparisonCard", "Конкретные отличия", "L6-3", "—", "P1", "⬜ Не начато"),
    ("L6-1", "M-K5", "Temporal", "UC: Temporal Diff + TimelineCard",
     "«Что менялось» по git-истории конфигурации",
     "git history parse", "Temporal UC, TimelineCard", "Timeline изменений", "M-K3", "—", "P2", "⬜ Не начато"),
    ("L6-4", "M-K5", "8.5-Ready 💰", "8.5-Ready Assessment (REVENUE)",
     "Standalone платный сервис: загрузка .dt → оплата → PDF-отчёт рисков перехода на 8.5",
     "landing + payment + автопайплайн + Diff 8.3↔8.5 API", "8.5-Ready сервис", "≥10 платных кейсов/мес, чек 9.9–49к₽", "L6-2,M-K4 compliance", "—", "P0", "⬜ Не начато"),
    ("DIST2", "M-K5", "Installer v2", "Electron installer v2.0.0",
     "~250МБ: EPF+CFE+Install-ps1 + auto-update + bundled JRE17 + BGE-M3 download-on-first-run (решение local-first)",
     "electron-builder, electron-updater, EV/OV cert", "analyst-setup-v2.0.0.exe", "Ставится на чистой Win10/11", "M-K3 EPF/CFE", "—", "P1", "⬜ Не начато"),
    ("OPS-3", "M-K5", "Export", "Export/Import .klzip", "Перенос knowledge-бандла между машинами", "zip bundle", ".klzip", "Bundle переносится", "M-K3", "—", "P2", "⬜ Не начато"),
    ("K5.99", "M-K5", "SUMMARY", "M-K5 SUMMARY", "Итог + pre-flight M-K6", "markdown", "M-K5/SUMMARY.md", "Фазы закрыты", "L6-4,DIST2", "—", "P1", "⬜ Не начато"),

    # ── M-K6 (крупно) ──
    ("L6-6", "M-K6", "NL2SQL", "UC: NL2SQL over config",
     "«Поставщики без оплаты за 90 дней» → запрос + TableCard",
     "NL→1С-запрос через graph+LLM", "NL2SQL UC", "Корректный запрос + карточка", "все M-K5", "—", "P1", "⬜ Не начато"),
    ("L6-7", "M-K6", "Code review", "UC: Auto Code Review",
     "Git webhook → комменты на PR с file:line + ИТС-ссылками",
     "GitHub/GitLab webhook, detector", "webhook bot", "Комментит PR", "M-K3 15.4", "—", "P1", "⬜ Не начато"),
    ("L6-5", "M-K6", "Benchmark", "UC: Cross-config benchmark",
     "Сравнение с отраслью при ≥50 проиндексированных конфигурациях (opt-in)",
     "анонимная агрегация fingerprints", "benchmark", "Работает при ≥50 конфигов", "L6-3", "—", "P2", "⬜ Не начато"),
    ("PG", "M-K6", "Pricing/GTM", "Pricing-tiers + франшиза + self-hosted",
     "Per-Configuration Tier + Franchise Bulk License + Self-Hosted (КИИ)",
     "billing, лицензирование", "pricing + контракты", "≥3 франчайзи, ≥1 self-hosted, MRR ≥500к₽", "—", "—", "P1", "⬜ Не начато"),
]


# ════════════════════════════════════════════════════════════════════════════
# Лист 3 — КАЛЕНДАРЬ И ЯКОРЯ
# ════════════════════════════════════════════════════════════════════════════
CAL_COLS = [
    ("Дата / окно", 18, True), ("Событие", 28, True), ("Тип", 12, True),
    ("Почему важно", 40, True), ("Что должно быть готово", 36, True), ("Связь", 14, True),
]
CALENDAR = [
    ("Лето 2026", "Переход клиентов УТ на 1С 8.5.x", "Внешний якорь", "Растёт спрос на 8.5-аудит", "8.5-Ready Assessment", "M-K5"),
    ("31.08.2026", "Дедлайн заявки на A&PM EVENT", "Внешний якорь", "Без заявки нет доклада/стенда осенью", "Подать заявку в 1-ю неделю M-K3", "M-K3"),
    ("01.10.2026", "Конец free-периода 1С:Напарника", "Внешний якорь", "Закрывается конкурентное окно; повод build-vs-buy", "OPS-2 telemetry собрана заранее", "M-K4"),
    ("Октябрь 2026", "INFOSTART A&PM EVENT", "Внешний якорь", "Ключевой канал лидов (цель ≥50 trial)", "MVP «Объяснитель» (M-K3) + доклад", "M-K3/K4"),
    ("Порядок →", "M-K3 → M-K4 → M-K5 → M-K6", "Внутр. последовательность", "Планируем усилием, не календарём (соло+AI сжимает сроки)", "Каждый милстоун атомарно полезен", "—"),
]
CAL_NOTE = "Внешние якоря = жёсткие (зависят от рынка/конференции). Внутренняя последовательность = по усилию и зависимостям, без привязки к календарным датам старых планов."


# ════════════════════════════════════════════════════════════════════════════
# Лист 4 — РИСКИ
# ════════════════════════════════════════════════════════════════════════════
RISK_COLS = [
    ("Риск", 34, True), ("Severity", 11, False), ("Триггер (где сработает)", 22, True),
    ("Mitigation", 44, True), ("Этап", 12, True),
]
RISKS = [
    ("TreeSitter BSL-грамматика неполна для редких конструкций", "HIGH", "M-K3.17.1 граф", "Fallback на текстовый парсинг; парсер целимся ≥95%, не 100%", "M-K3"),
    ("MetaVision headless не собирается за 3 дня", "MEDIUM", "M-K4.16.0 spike", "Hard-timebox 3 дн → наш GraphCard становится primary (6 дн вместо 15)", "M-K4"),
    ("Размер bundle (BGE-M3 570МБ + JRE + BSL LS jar)", "HIGH", "M-K5 installer v2", "Решено: download-on-first-run, НЕ упаковка; jlink+proguard для JRE/jar", "M-K5"),
    ("Citation-контракт G1 инвазивен для loop.py", "MEDIUM", "M-K4 citation layer", "Вводить за feature-flag, не ломать текущий формат ответа", "M-K4"),
    ("1С:Напарник станет дешёвым/бесплатным", "HIGH", "после 01.10.2026", "OPS-2 telemetry → build-vs-buy; наша ниша L4–L6 (diagnose), не writing", "M-K4/K5"),
    ("Юр.вопрос по reference-конфигурациям типовых", "MEDIUM", "M-K5 reference lib", "Только fingerprint + структура, НЕ код типовой", "M-K5"),
    ("Diagnose-правилам нужны реальные кейсы", "MEDIUM", "M-K3.17.2", "Собрать ≥20 кейсов из практики Транзит/УСО ДО старта (задача K3.0b)", "M-K3"),
    ("Рассинхрон планов/STATE копится", "LOW", "постоянно", "Действие D0 + единый STATE после каждой фазы", "Гигиена"),
    ("БГУ/ЗУП конфликтуют с CFE", "MEDIUM", "M-K3.13b", "feature-detection per config (13b.0); 5 раздельных smoke-наборов", "M-K3"),
    ("Правило 3 итераций (Матаков)", "—", "любая AI-задача", "За 3 раунда без прогресса → STOP, переформулировать spec / писать руками", "Все"),
]


# ════════════════════════════════════════════════════════════════════════════
# Лист 5 — ПАРАЛЛЕЛЬНЫЕ ТРЕКИ И ДОЛГ
# ════════════════════════════════════════════════════════════════════════════
TRACK_COLS = [
    ("Трек", 26, True), ("Что это", 40, True), ("Статус", 22, True),
    ("Блокирует основную работу?", 22, True), ("Действие", 36, True),
]
TRACKS = [
    ("M-K2.5 NIM rebuild карточек", "63 203 эталонных карточки типовых через qwen-122b (фоновый процесс)", "🟡 ~6.5% (2026-05-28)", "НЕТ — граф/EPF/CFE от карточек не зависят", "D3: проверить жив ли процесс; держать как enrichment"),
    ("Гигиена релиза оболочки", "Код v1.4.8, последний тег v1.2.2 — сборки 1.3/1.4 не тегались", "🟡 долг", "НЕТ", "Smoke на VM + EV/OV cert + git tag v1.4.x + 3 Playwright spec"),
    ("Технический долг", "POST-RELEASE-DEBT (14), 6 cards refactor (1/6), 7 flaky-тестов (cp1251), F841/E501 в loop.py", "🟡 backlog", "НЕТ", "По 1–2 пункта/нед как гигиенические коммиты"),
    ("71 MEDIUM/LOW finding аудита", "Остаток общего аудита (99 findings, 28 закрыты в M-K0)", "🟡 backlog", "НЕТ", "Приоритезация владельцем, параллельно с M-K1..K6"),
]


# ════════════════════════════════════════════════════════════════════════════
# Лист 6 — ПРИНЦИПЫ И РЕШЕНИЯ
# ════════════════════════════════════════════════════════════════════════════
PRIN_COLS = [
    ("#", 5, False), ("Принцип / Решение", 30, True), ("Суть", 60, True), ("Тип", 14, True),
]
PRINCIPLES = [
    ("1", "Local-first", "Индексация + embeddings локально (FastEmbed ONNX). Код клиента не уходит в облако без opt-in. Compliance 152-ФЗ/КИИ", "Принцип"),
    ("2", "Grounded", "Любой ответ Knowledge Layer ссылается на file:line + объект метаданных + стандарт ИТС/БСП. Без grounding = баг", "Принцип"),
    ("3", "Per-config isolation", "Каждая база клиента — отдельный индекс ~/.analyst-1c/knowledge/<fingerprint>/. Никаких глобальных индексов между конфигурациями", "Принцип"),
    ("4", "Incremental", "mtime-watcher + diff-based update; полная переиндексация только при fingerprint mismatch", "Принцип"),
    ("5", "Test-first для каждого UC", "UC = запрос + ожидаемый ответ. Реализация начинается с golden-кейса, не с кода", "Принцип"),
    ("6", "Trace везде", "Каждый ответ = развёрнутый ToolTrace (workflow, источники, промежуточные результаты)", "Принцип"),
    ("7", "Open-core boundary", "Open: pipeline/AST/graph-схема/MCP-скелет. Closed: reference-типовые, business rulebook, knowledge-бандлы", "Принцип"),
    ("D", "Embeddings = local BGE-M3", "РЕШЕНО 2026-05-29: FastEmbed/BGE-M3 (ADR-003). Ключ не нужен; +570МБ → download-on-first-run в M-K5. OpenAI — опц. fallback", "Решение"),
    ("D", "CFE-именование", "Подсистема АналитикПлюс + префикс АП_ для общих модулей (Q1)", "Решение"),
    ("D", "Типовые для CFE", "УТ + ERP + КА + БГУ + ЗУП (5 на БСП 3.1+) (Q2)", "Решение"),
    ("D", "EPF раньше CFE", "Phase 13a (EPF) до Phase 13b (CFE) (Q3)", "Решение"),
    ("D", "BSL LS = detector-first", "Detector-режим в M-K3, streaming опционально в M-K4 (Q5/G5)", "Решение"),
    ("AG", "Anti-goal: не свой LLM", "Используем NVIDIA NIM / Cloud.ru / DeepSeek", "Anti-goal"),
    ("AG", "Anti-goal: не редактор кода", "Мы analyzer/explainer, не IDE и не пишем код за пользователя (это Напарник)", "Anti-goal"),
    ("AG", "Anti-goal: не cloud-индексация чужих баз", "Local-first до M-K6; не Neo4j/Kuzu (SQLite CTE достаточно); не proprietary embeddings", "Anti-goal"),
]


def main():
    wb = Workbook()
    wb.remove(wb.active)

    add_sheet(wb, "Обзор", "1С АНАЛИТИК · Обзор милстоунов (2026-05-29)", OVERVIEW_COLS, OVERVIEW, status_col=8)
    # примечание под обзором
    ws = wb["Обзор"]
    note_row = 2 + 1 + len(OVERVIEW) + 1
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=len(OVERVIEW_COLS))
    nc = ws.cell(row=note_row, column=1, value=OVERVIEW_NOTE)
    nc.font = Font(name="Calibri", italic=True, size=9, color="6B6B6B")
    nc.alignment = WRAP_TOP
    ws.row_dimensions[note_row].height = 42

    add_sheet(wb, "Workflow", "WORKFLOW · Что и КАК делается (детально M-K3/M-K4, крупно M-K5/M-K6)",
              WF_COLS, WORKFLOW, status_col=12, prio_col=11)
    add_sheet(wb, "Календарь и якоря", "КАЛЕНДАРЬ И ВНЕШНИЕ ЯКОРЯ", CAL_COLS, CALENDAR)
    ws = wb["Календарь и якоря"]
    nr = 2 + 1 + len(CALENDAR) + 1
    ws.merge_cells(start_row=nr, start_column=1, end_row=nr, end_column=len(CAL_COLS))
    nc = ws.cell(row=nr, column=1, value=CAL_NOTE)
    nc.font = Font(name="Calibri", italic=True, size=9, color="6B6B6B")
    nc.alignment = WRAP_TOP
    ws.row_dimensions[nr].height = 36

    add_sheet(wb, "Риски", "РИСК-РЕЕСТР", RISK_COLS, RISKS, status_col=2)
    add_sheet(wb, "Параллельные треки", "ПАРАЛЛЕЛЬНЫЕ ТРЕКИ И ТЕХ-ДОЛГ", TRACK_COLS, TRACKS, status_col=3)
    add_sheet(wb, "Принципы и решения", "ПРИНЦИПЫ · РЕШЕНИЯ · ANTI-GOALS", PRIN_COLS, PRINCIPLES)

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Workflow_1C_Analyst_2026-05-29.xlsx")
    wb.save(out)
    total_days = sum(
        float(r[9]) for r in WORKFLOW
        if str(r[9]).replace(".", "", 1).isdigit()
    )
    print(f"OK: {out}")
    print(f"Листов: {len(wb.sheetnames)} | задач в Workflow: {len(WORKFLOW)} | "
          f"оценённое усилие (числовые дн): ~{total_days:.0f} человеко-дней")


if __name__ == "__main__":
    main()
